"""One requested operation per process. No background recording or shell tool."""
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys


def execute(request):
    kind = request["kind"]
    data = request.get("input", {})
    if kind == "status":
        return {name: importlib.util.find_spec(name) is not None
                for name in ("docx", "faster_whisper", "pyautogui")}
    if kind == "document.word":
        from docx import Document
        document = Document()
        document.add_heading(data["title"], 0)
        for line in data["body"].splitlines():
            document.add_paragraph(line)
        document.save(data["output"])
        return {"created": True}
    if kind == "voice.transcribe":
        from faster_whisper import WhisperModel
        # A local directory is required. A request cannot choose a remote model.
        model_path = Path(data["model"])
        if not (model_path / "model.bin").is_file():
            raise ValueError("VOICE_MODEL_NOT_INSTALLED")
        import av
        import numpy as np
        # Bound decoded audio as well as the upload size. Compressed files must
        # not expand into hours of audio in a single companion process.
        chunks, samples = [], 0
        with av.open(data["audio"]) as source:
            resampler = av.AudioResampler(format="s16", layout="mono", rate=16000)
            for frame in source.decode(audio=0):
                for converted in resampler.resample(frame):
                    values = converted.to_ndarray().flatten()
                    samples += values.size
                    if samples > 16000 * 60:
                        raise ValueError("AUDIO_DURATION_TOO_LONG")
                    chunks.append(values)
        if not chunks:
            raise ValueError("AUDIO_EMPTY")
        audio = np.concatenate(chunks).astype(np.float32) / 32768.0
        model = WhisperModel(str(model_path), device="cpu", compute_type="int8", local_files_only=True, cpu_threads=2, num_workers=1)
        segments, info = model.transcribe(audio, beam_size=1, vad_filter=True)
        text = " ".join(s.text.strip() for s in segments).strip()
        return {"text": text[:12000], "language": info.language}
    if kind == "desktop.action":
        import pyautogui
        pyautogui.FAILSAFE = True
        if data["action"] == "open":
            apps = {"notepad": "notepad.exe", "calculator": "calc.exe"}
            if data.get("app") not in apps:
                raise ValueError("DESKTOP_APP_DENIED")
            subprocess.Popen([apps[data["app"]]], shell=False)
            return {"launched": data["app"], "verifiedWindow": False}
        target = data.get("targetTitle")
        matches = [window for window in pyautogui.getWindowsWithTitle(target or "") if window.title == target]
        if len(matches) != 1:
            raise ValueError("DESKTOP_WINDOW_CHANGED")
        matches[0].activate()
        title = pyautogui.getActiveWindowTitle() or ""
        if not data.get("targetTitle") or title != data["targetTitle"]:
            raise ValueError("DESKTOP_WINDOW_CHANGED")
        if data["action"] == "type":
            text = data.get("text", "")
            if not isinstance(text, str) or not text.isascii() or len(text) > 1000 or any(ord(c)<32 for c in text):
                raise ValueError("DESKTOP_TEXT_INVALID")
            pyautogui.write(text, interval=0.015)
            return {"typedCharacters": len(text), "window": title}
        if data["action"] == "click":
            x, y = data["x"], data["y"]
            width, height = pyautogui.size()
            if not isinstance(x, int) or not isinstance(y, int) or not (0 <= x < width and 0 <= y < height):
                raise ValueError("DESKTOP_POSITION_INVALID")
            pyautogui.click(x=x, y=y)
            return {"clicked": {"x": x, "y": y}, "window": title}
        raise ValueError("DESKTOP_ACTION_DENIED")
    raise ValueError("PYTHON_OPERATION_DENIED")


if __name__ == "__main__":
    try:
        result = execute(json.load(sys.stdin))
        print(json.dumps({"ok": True, "result": result}, ensure_ascii=True))
    except Exception as exc:
        # Do not echo request text, paths, credentials or raw provider responses.
        code = str(exc) if isinstance(exc, ValueError) else type(exc).__name__
        print(json.dumps({"ok": False, "error": code[:100]}))
        sys.exit(1)
